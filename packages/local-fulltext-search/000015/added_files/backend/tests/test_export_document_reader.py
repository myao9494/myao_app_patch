"""
AIインプット書き出しの文書読み込み（app.vector.export_document_reader）のテスト。

仕様:
- Markdown / テキストは文字コードを判定して読み、.docx / .pptx は本文テキストを抽出して読む。
- Excel・PDF・画像・その他の拡張子は変換対象外（None）。バイナリを Markdown 変換へ流さない。
- HTML 書き出しと Markdown+画像の書き出しは、対象外の文書を出力に含めない。
"""

from __future__ import annotations

from pathlib import Path

from app.vector.export_document_reader import read_export_document_text
from app.vector.html_exporter import export_documents_to_html
from app.vector.markdown_image_exporter import export_documents_to_markdown_and_image


def _make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("ワード本文のテキスト")
    doc.save(str(path))


def _make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "パワポ本文のテキスト"
    prs.save(str(path))


def test_markdown_and_text_are_read(tmp_path: Path) -> None:
    md = tmp_path / "a.md"
    md.write_text("# 見出し", encoding="utf-8")
    txt = tmp_path / "b.txt"
    txt.write_bytes("日本語".encode("cp932"))
    assert read_export_document_text(md) == "# 見出し"
    assert read_export_document_text(txt) == "日本語"


def test_docx_and_pptx_are_extracted(tmp_path: Path) -> None:
    docx_path = tmp_path / "a.docx"
    pptx_path = tmp_path / "b.pptx"
    _make_docx(docx_path)
    _make_pptx(pptx_path)
    assert "ワード本文のテキスト" in (read_export_document_text(docx_path) or "")
    assert "パワポ本文のテキスト" in (read_export_document_text(pptx_path) or "")


def test_broken_office_file_becomes_empty_text(tmp_path: Path) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    assert read_export_document_text(broken) == ""


def test_other_extensions_are_not_converted(tmp_path: Path) -> None:
    for name in ("t.xlsx", "t.xlsm", "t.pdf", "t.png", "t.msg", "t.zip", "noext"):
        p = tmp_path / name
        p.write_bytes(b"\x00\x01binary")
        assert read_export_document_text(p) is None, name


def test_html_export_skips_unsupported_and_reads_office(tmp_path: Path) -> None:
    md = tmp_path / "note.md"
    md.write_text("# ノート本文", encoding="utf-8")
    docx_path = tmp_path / "doc.docx"
    pptx_path = tmp_path / "slide.pptx"
    _make_docx(docx_path)
    _make_pptx(pptx_path)
    xlsx = tmp_path / "sheet.xlsx"
    xlsx.write_bytes(b"PK\x03\x04binary")

    html_content, stats = export_documents_to_html(
        file_paths=[str(md), str(docx_path), str(pptx_path), str(xlsx)],
    )

    assert stats["total_documents"] == 3
    assert "ワード本文のテキスト" in html_content
    assert "パワポ本文のテキスト" in html_content
    assert "sheet.xlsx" not in html_content


def test_markdown_export_skips_unsupported_and_reads_office(tmp_path: Path) -> None:
    md = tmp_path / "note.md"
    md.write_text("# ノート本文", encoding="utf-8")
    pptx_path = tmp_path / "slide.pptx"
    _make_pptx(pptx_path)
    xlsx = tmp_path / "sheet.xlsx"
    xlsx.write_bytes(b"PK\x03\x04binary")

    text, _image, stats = export_documents_to_markdown_and_image(
        file_paths=[str(md), str(pptx_path), str(xlsx)],
        vault_path=str(tmp_path),
    )

    assert "パワポ本文のテキスト" in text
    assert "sheet.xlsx" not in text
    assert stats["total_documents"] == 2
